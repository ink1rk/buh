"""Импорт CSV/XLSX: загрузка → маппинг колонок → проверка → применение.

Первичное наполнение CMDB обычно начинается с чьей-то таблицы, поэтому импорт
должен прощать русские заголовки, точку с запятой в качестве разделителя и cp1251.
"""

from __future__ import annotations

import csv
import io
import uuid
from datetime import UTC, datetime
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from itms.core.context import ActorKind, RequestContext, current_context, use_context
from itms.core.errors import Invalid, NotFound
from itms.models.cmdb import Ci, Location
from itms.models.directory import Employee
from itms.models.enums import (
    CiStatus,
    CiType,
    Criticality,
    Environment,
    ImportStatus,
    ImportTarget,
)
from itms.models.system import ImportJob
from itms.services import ci_service, directory_service, location_service

PREVIEW_ROWS = 20

#: Русские и английские заголовки, которые распознаются автоматически.
FIELD_ALIASES: dict[ImportTarget, dict[str, tuple[str, ...]]] = {
    ImportTarget.CI: {
        "code": ("код", "code", "шифр", "идентификатор"),
        "name": ("наименование", "название", "имя", "name", "объект"),
        "ci_type": ("тип", "type", "тип объекта"),
        "status": ("статус", "состояние", "status"),
        "criticality": ("критичность", "criticality"),
        "environment": ("среда", "контур", "environment"),
        "location": ("размещение", "местоположение", "локация", "помещение", "location"),
        "owner": ("ответственный", "владелец", "owner"),
        "vendor": ("производитель", "вендор", "vendor", "бренд"),
        "model": ("модель", "model"),
        "serial_number": ("серийный номер", "серийник", "s/n", "sn", "serial"),
        "inventory_number": ("инвентарный номер", "инв. номер", "инвентарник", "inventory"),
        "description": ("описание", "комментарий", "примечание", "description"),
        "tags": ("теги", "метки", "tags"),
    },
    ImportTarget.LOCATION: {
        "name": ("наименование", "название", "имя", "name"),
        "location_type": ("тип", "type"),
        "parent": ("родитель", "входит в", "parent"),
        "code": ("код", "code"),
        "address": ("адрес", "address"),
        "description": ("описание", "description"),
    },
    ImportTarget.EMPLOYEE: {
        "full_name": ("фио", "сотрудник", "имя", "full_name", "name"),
        "position": ("должность", "position"),
        "email": ("email", "почта", "эл. почта"),
        "phone": ("телефон", "phone"),
        "telegram": ("telegram", "телеграм"),
        "department": ("отдел", "подразделение", "department"),
        "support_line": ("линия", "линия поддержки", "support_line"),
    },
}

REQUIRED_FIELDS: dict[ImportTarget, tuple[str, ...]] = {
    ImportTarget.CI: ("name", "ci_type"),
    ImportTarget.LOCATION: ("name", "location_type"),
    ImportTarget.EMPLOYEE: ("full_name",),
}


def parse_table(filename: str, content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    lowered = filename.lower()
    if lowered.endswith(".xlsx") or lowered.endswith(".xlsm"):
        return _parse_xlsx(content)
    if lowered.endswith(".csv") or lowered.endswith(".txt"):
        return _parse_csv(content)
    raise Invalid("Поддерживаются файлы CSV и XLSX", code_hint="unsupported_file")


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise Invalid("Не удалось определить кодировку файла", code_hint="bad_encoding")


def _parse_csv(content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    text = _decode(content)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    columns = [c.strip() for c in (reader.fieldnames or []) if c and c.strip()]
    rows = [
        {(k or "").strip(): (v or "").strip() for k, v in row.items() if k}
        for row in reader
    ]
    return columns, [r for r in rows if any(r.values())]


def _parse_xlsx(content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    try:
        header = next(iterator)
    except StopIteration as exc:
        raise Invalid("Файл пуст", code_hint="empty_file") from exc
    columns = [str(c).strip() for c in header if c is not None and str(c).strip()]
    rows: list[dict[str, str]] = []
    for raw in iterator:
        values = ["" if v is None else str(v).strip() for v in raw]
        if not any(values):
            continue
        rows.append({col: (values[i] if i < len(values) else "") for i, col in enumerate(columns)})
    workbook.close()
    return columns, rows


def suggest_mapping(target: ImportTarget, columns: list[str]) -> dict[str, str]:
    """Подбирает соответствие «колонка файла → поле сущности» по названию заголовка."""
    aliases = FIELD_ALIASES[target]
    mapping: dict[str, str] = {}
    used: set[str] = set()
    for column in columns:
        key = column.strip().lower()
        for field_name, variants in aliases.items():
            if field_name in used:
                continue
            if key == field_name or key in variants:
                mapping[column] = field_name
                used.add(field_name)
                break
    return mapping


async def create_job(
    session: AsyncSession, *, target: ImportTarget, filename: str, content: bytes
) -> ImportJob:
    columns, rows = parse_table(filename, content)
    if not columns:
        raise Invalid("В файле не найдены заголовки колонок", code_hint="no_columns")
    job = ImportJob(
        target=target,
        status=ImportStatus.DRAFT,
        filename=filename,
        columns=columns,
        mapping=suggest_mapping(target, columns),
        rows_total=len(rows),
        preview=rows[:PREVIEW_ROWS],
        options={"raw_rows": rows},
        created_by=current_context().actor_id,
    )
    session.add(job)
    await session.flush()
    return job


async def get_job(session: AsyncSession, job_id: uuid.UUID) -> ImportJob:
    job = (
        await session.execute(select(ImportJob).where(ImportJob.id == job_id))
    ).scalar_one_or_none()
    if job is None:
        raise NotFound("Задание импорта не найдено", entity_id=str(job_id))
    return job


def _mapped_rows(job: ImportJob) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = job.options.get("raw_rows", [])
    mapping: dict[str, str] = job.mapping or {}
    result = []
    for row in rows:
        mapped = {field: (row.get(column) or "").strip() for column, field in mapping.items()}
        result.append(mapped)
    return result


def _enum_or_error(
    enum_cls: type, value: str, field_name: str, errors: list[str], default: Any = None
) -> Any:
    if not value:
        return default
    try:
        return enum_cls(value.strip().upper())
    except ValueError:
        errors.append(f"{field_name}: недопустимое значение «{value}»")
        return default


async def validate_job(session: AsyncSession, job: ImportJob) -> ImportJob:
    missing = [f for f in REQUIRED_FIELDS[job.target] if f not in set((job.mapping or {}).values())]
    if missing:
        raise Invalid(
            "Не сопоставлены обязательные поля",
            code_hint="mapping_incomplete",
            missing=missing,
        )

    rows = _mapped_rows(job)
    errors: list[dict[str, Any]] = []
    valid = 0
    for index, row in enumerate(rows, start=2):  # строка 1 — заголовок
        row_errors: list[str] = []
        for field_name in REQUIRED_FIELDS[job.target]:
            if not row.get(field_name):
                row_errors.append(f"{field_name}: значение обязательно")
        if job.target == ImportTarget.CI:
            _enum_or_error(CiType, row.get("ci_type", ""), "ci_type", row_errors)
            _enum_or_error(CiStatus, row.get("status", ""), "status", row_errors)
            _enum_or_error(Criticality, row.get("criticality", ""), "criticality", row_errors)
            _enum_or_error(Environment, row.get("environment", ""), "environment", row_errors)
            if row.get("location"):
                found = await _find_location(session, row["location"])
                if found is None:
                    row_errors.append(f"location: размещение «{row['location']}» не найдено")
        if row_errors:
            errors.append({"row": index, "errors": row_errors})
        else:
            valid += 1

    job.rows_valid = valid
    job.rows_invalid = len(rows) - valid
    job.errors = errors[:200]
    job.status = ImportStatus.VALIDATED if valid else ImportStatus.FAILED
    await session.flush()
    return job


async def _find_location(session: AsyncSession, value: str) -> Location | None:
    value = value.strip()
    return (
        await session.execute(
            select(Location)
            .where((Location.path == value) | (Location.name == value) | (Location.code == value))
            .limit(1)
        )
    ).scalar_one_or_none()


async def _find_employee(session: AsyncSession, value: str) -> Employee | None:
    value = value.strip()
    return (
        await session.execute(
            select(Employee)
            .where((Employee.full_name == value) | (Employee.email == value))
            .limit(1)
        )
    ).scalar_one_or_none()


async def apply_job(session: AsyncSession, job: ImportJob) -> ImportJob:
    if job.status not in (ImportStatus.VALIDATED, ImportStatus.DRAFT):
        raise Invalid("Задание импорта уже применено", code_hint="import_already_applied")
    if job.status == ImportStatus.DRAFT:
        await validate_job(session, job)

    base = current_context()
    import_ctx = RequestContext(
        request_id=base.request_id,
        actor_id=base.actor_id,
        actor_kind=ActorKind.IMPORT,
        actor_label=f"Импорт: {job.filename}",
        source="import",
        provenance=base.provenance,
    )

    created = updated = 0
    with use_context(import_ctx):
        for row in _mapped_rows(job):
            try:
                if job.target == ImportTarget.CI:
                    was_created = await _apply_ci_row(session, row)
                elif job.target == ImportTarget.LOCATION:
                    was_created = await _apply_location_row(session, row)
                else:
                    was_created = await _apply_employee_row(session, row)
            except (Invalid, NotFound):
                continue
            created += int(was_created)
            updated += int(not was_created)

    job.rows_created = created
    job.rows_updated = updated
    job.status = ImportStatus.APPLIED
    job.applied_at = datetime.now(UTC)
    await session.flush()
    return job


async def _apply_ci_row(session: AsyncSession, row: dict[str, str]) -> bool:
    errors: list[str] = []
    ci_type = _enum_or_error(CiType, row.get("ci_type", ""), "ci_type", errors, CiType.OTHER)
    if errors or not row.get("name"):
        raise Invalid("Строка не прошла проверку")

    location = await _find_location(session, row["location"]) if row.get("location") else None
    owner = await _find_employee(session, row["owner"]) if row.get("owner") else None
    payload: dict[str, Any] = {
        "ci_type": ci_type,
        "name": row["name"],
        "code": row.get("code") or None,
        "vendor": row.get("vendor") or None,
        "model": row.get("model") or None,
        "serial_number": row.get("serial_number") or None,
        "inventory_number": row.get("inventory_number") or None,
        "description": row.get("description") or None,
        "location_id": location.id if location else None,
        "owner_employee_id": owner.id if owner else None,
    }
    status = _enum_or_error(CiStatus, row.get("status", ""), "status", errors)
    criticality = _enum_or_error(Criticality, row.get("criticality", ""), "criticality", errors)
    environment = _enum_or_error(Environment, row.get("environment", ""), "environment", errors)
    if criticality:
        payload["criticality"] = criticality
    if environment:
        payload["environment"] = environment
    if row.get("tags"):
        payload["tags"] = [t.strip() for t in row["tags"].replace(";", ",").split(",") if t.strip()]

    existing = None
    if payload["code"]:
        existing = (
            await session.execute(select(Ci).where(Ci.code == payload["code"]))
        ).scalar_one_or_none()
    if existing is None and payload["serial_number"]:
        existing = (
            await session.execute(
                select(Ci).where(Ci.serial_number == payload["serial_number"])
            )
        ).scalar_one_or_none()

    if existing is None:
        if status:
            payload["status"] = status
        await ci_service.create_ci(session, payload)
        return True

    payload.pop("ci_type", None)
    if status:
        payload["status"] = status
    await ci_service.update_ci(session, existing.id, payload)
    return False


async def _apply_location_row(session: AsyncSession, row: dict[str, str]) -> bool:
    from itms.models.enums import LocationType

    errors: list[str] = []
    location_type = _enum_or_error(
        LocationType, row.get("location_type", ""), "location_type", errors
    )
    if errors or location_type is None or not row.get("name"):
        raise Invalid("Строка не прошла проверку")
    parent = await _find_location(session, row["parent"]) if row.get("parent") else None
    existing = await _find_location(session, row["name"])
    if existing is not None:
        await location_service.update_location(
            session,
            existing.id,
            {"code": row.get("code") or None, "address": row.get("address") or None},
        )
        return False
    await location_service.create_location(
        session,
        {
            "name": row["name"],
            "location_type": location_type,
            "parent_id": parent.id if parent else None,
            "code": row.get("code") or None,
            "address": row.get("address") or None,
            "description": row.get("description") or None,
        },
    )
    return True


async def _apply_employee_row(session: AsyncSession, row: dict[str, str]) -> bool:
    if not row.get("full_name"):
        raise Invalid("Строка не прошла проверку")
    existing = await _find_employee(session, row["full_name"])
    payload = {
        "full_name": row["full_name"],
        "position": row.get("position") or None,
        "email": row.get("email") or None,
        "phone": row.get("phone") or None,
        "telegram": row.get("telegram") or None,
    }
    if existing is not None:
        await directory_service.update_employee(session, existing.id, payload)
        return False
    await directory_service.create_employee(session, payload)
    return True
