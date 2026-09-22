"""Tests for the bank connector layer and the Ozon Bank provider."""

from __future__ import annotations

import io
from pathlib import Path
from datetime import date

import httpx
import pytest
from httpx import AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.connectors import get_connector
from app.connectors.base import (
    RawOperation,
    StatementParseError,
    parse_money,
    parse_statement_date,
)
from app.connectors.categorize import classify
from app.connectors.open_banking import (
    OpenBankingClient,
    OpenBankingError,
    statement_from_transactions,
)
from app.connectors.tabular import read_csv_rows, rows_to_statement
from app.models.account import Account
from app.models.connection import BankConnection, BankOperation
from app.models.transaction import Transaction
from app.services.bank_sync import ingest_statement, refresh_fingerprints

OZON_CSV = """Выписка по счёту Ozon Банк
Клиент: Кирилл
Период: 01.03.2026 - 31.03.2026

Дата операции;Дата обработки;Сумма операции в валюте счёта;Валюта;Категория;MCC;Описание операции;Статус;Остаток после операции
12.03.2026 14:23;12.03.2026;-1 234,56 ₽;RUB;Супермаркеты;5411;Пятёрочка;Проведена;48 765,44
13.03.2026 09:01;13.03.2026;+180 000,00 ₽;RUB;Зарплата;;Зарплата ООО Ромашка;Проведена;228 765,44
14.03.2026 10:15;14.03.2026;−450,00 ₽;RUB;Кофейни;5814;Surf Coffee;Проведена;228 315,44
15.03.2026 11:00;15.03.2026;-2 000,00 ₽;RUB;Переводы;;Перевод между своими счетами;Проведена;226 315,44
16.03.2026 12:00;16.03.2026;-999,00 ₽;RUB;Подписки;5815;Netflix;Отклонена;226 315,44
17.03.2026 13:00;17.03.2026;+320,50 ₽;RUB;Кэшбэк;;Кэшбэк за март;Проведена;226 635,94
"""


def _csv_bytes(text: str = OZON_CSV, encoding: str = "utf-8") -> bytes:
    return text.encode(encoding)


# --- primitives -------------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("-1 234,56 ₽", -1234.56),
        ("\u2212450,00", -450.0),
        ("+180 000,00 ₽", 180000.0),
        ("1\u00a0234,56", 1234.56),
        ("1.234,56", 1234.56),
        ("1,234.56", 1234.56),
        ("(2 000,00)", -2000.0),
        ("1234", 1234.0),
        ("", None),
        ("—", None),
        (-55.5, -55.5),
    ],
)
def test_parse_money(raw, expected):
    assert parse_money(raw) == expected


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("12.03.2026", date(2026, 3, 12)),
        ("12.03.2026 14:23", date(2026, 3, 12)),
        ("2026-03-12T10:00:00", date(2026, 3, 12)),
        ("12/03/2026", date(2026, 3, 12)),
        ("12.03.26", date(2026, 3, 12)),
        ("5 марта 2026", date(2026, 3, 5)),
        ("32.13.2026", None),
        ("", None),
    ],
)
def test_parse_statement_date(raw, expected):
    assert parse_statement_date(raw) == expected


# --- statement parsing ------------------------------------------------------


def test_parses_ozon_csv_with_preamble_and_unicode_minus():
    statement = get_connector("ozon").parse_statement(_csv_bytes(), "statement.csv")

    # The declined Netflix row must not reach the ledger.
    assert len(statement.operations) == 5
    amounts = [op.amount for op in statement.operations]
    assert amounts == [-1234.56, 180000.0, -450.0, -2000.0, 320.5]
    assert statement.period_from == date(2026, 3, 12)
    assert statement.period_to == date(2026, 3, 17)
    assert statement.closing_balance == 226635.94

    groceries = statement.operations[0]
    assert groceries.mcc == "5411"
    assert groceries.bank_category == "Супермаркеты"
    assert groceries.merchant == "Пятёрочка"
    assert groceries.currency == "RUB"


def test_parses_cp1251_csv():
    # CP1251 has no ₽ or unicode minus, so use the ASCII-signed variant.
    cp1251_csv = OZON_CSV.replace(" ₽", "").replace("\u2212", "-")
    statement = get_connector("ozon").parse_statement(
        _csv_bytes(cp1251_csv, encoding="cp1251"), "s.csv"
    )
    assert len(statement.operations) == 5
    assert statement.operations[0].merchant == "Пятёрочка"


def test_parses_utf8_bom_csv():
    statement = get_connector("ozon").parse_statement(
        b"\xef\xbb\xbf" + _csv_bytes(), "s.csv"
    )
    assert len(statement.operations) == 5


def test_parses_debit_credit_column_pair():
    csv = (
        "Дата;Приход;Расход;Описание\n"
        "12.03.2026;;1 234,56;Пятёрочка\n"
        "13.03.2026;180 000,00;;Зарплата\n"
    )
    statement = rows_to_statement(read_csv_rows(csv.encode()))
    assert [op.amount for op in statement.operations] == [-1234.56, 180000.0]


def test_uses_direction_column_when_amount_is_unsigned():
    csv = (
        "Дата операции;Сумма;Тип операции;Описание\n"
        "12.03.2026;1234,56;Списание;Пятёрочка\n"
        "13.03.2026;180000;Поступление;Зарплата\n"
    )
    statement = rows_to_statement(read_csv_rows(csv.encode()))
    assert [op.amount for op in statement.operations] == [-1234.56, 180000.0]


def test_ozon_bonus_points_are_not_money():
    csv = (
        "Дата операции;Сумма операции;Валюта;Описание операции\n"
        "14.03.2026;-500,00;балл;Списание баллов Ozon\n"
        "14.03.2026;-300,00;RUB;Пятёрочка\n"
    )
    statement = get_connector("ozon").parse_statement(csv.encode(), "s.csv")
    assert [op.amount for op in statement.operations] == [-300.0]
    assert any("баллах" in warning for warning in statement.warnings)


def test_statement_without_recognisable_header_is_rejected():
    with pytest.raises(StatementParseError, match="заголовка") as refusal:
        get_connector("ozon").parse_statement(b"just;some;text\n1;2;3\n", "s.csv")
    # По отказу должно быть видно, что именно прочитали: иначе спорить с ним
    # можно только вслепую.
    assert "Строк в файле: 2" in str(refusal.value)


def test_a_statement_without_a_header_is_read_by_its_rows():
    """Справка о движении средств приходит таблицей без шапки.

    Строк, похожих на операции, в ней сотни — по ним колонки и видно. Отказ
    «проверьте, что это выписка» доставался файлу, который выпиской и был.
    """
    rows = "".join(
        f"1{day:02d};0{day % 9 + 1}.03.2026;Оплата в магазине у дома;-{day}00,50\n"
        for day in range(1, 21)
    )
    statement = get_connector("ozon").parse_statement(rows.encode(), "spravka.csv")

    assert len(statement.operations) == 20
    assert statement.operations[0].amount == -100.5
    assert statement.operations[0].description == "Оплата в магазине у дома"
    assert any("по содержимому" in warning for warning in statement.warnings)


def test_a_couple_of_dated_lines_is_not_a_statement():
    """Счёт или квитанция — тоже строки с датой и суммой, но не история трат."""
    invoice = (
        "Счёт на оплату №12 от 01.03.2026;;\n"
        "Услуга;01.03.2026;12 000,00\n"
        "Доставка;01.03.2026;500,00\n"
    )
    with pytest.raises(StatementParseError, match="справка, счёт или квитанция"):
        get_connector("ozon").parse_statement(invoice.encode(), "invoice.csv")


def test_unsupported_provider_is_rejected():
    with pytest.raises(StatementParseError, match="не поддерживается"):
        get_connector("sberbank")


def test_parses_xlsx_statement():
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Выписка Ozon Банк"])
    sheet.append([])
    sheet.append(["Дата операции", "Сумма операции", "Категория", "MCC", "Описание операции"])
    sheet.append([date(2026, 3, 12), -1234.56, "Супермаркеты", "5411", "Пятёрочка"])
    sheet.append([date(2026, 3, 13), 180000, "Зарплата", None, "Зарплата"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    statement = get_connector("ozon").parse_statement(buffer.getvalue(), "statement.xlsx")
    assert [op.amount for op in statement.operations] == [-1234.56, 180000.0]
    assert statement.operations[0].mcc == "5411"


def _paged_xlsx(pages: list[list[list[object]]]) -> bytes:
    """Выписка, переведённая из PDF: лист на страницу и ложный размер листа.

    Так выглядит годовая выписка Ozon Банка: 288 листов по странице, шапка
    повторена на каждой, а в описании листа стоит `A1` — размер, которому
    нельзя верить.
    """
    import re
    import zipfile

    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for index, rows in enumerate(pages, 1):
        sheet = workbook.create_sheet(f"Sheet{index}")
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)

    spoiled = io.BytesIO()
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(spoiled, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename.startswith("xl/worksheets/"):
                data = re.sub(rb'<dimension ref="[^"]+"\s*/>', b'<dimension ref="A1"/>', data)
            target.writestr(item, data)
    return spoiled.getvalue()


OZON_HEADER = ["Дата операции", "Документ", "Назначение платежа", "Сумма операции"]


def _ozon_pages() -> list[list[list[object]]]:
    return [
        [
            ["ООО «ОЗОН Банк»"],
            ["Входящий остаток: 142.30 ₽"],
            OZON_HEADER,
            [None, None, None, "Российские рубли", "Валюта"],
            ["21.09.2026 18:51:01", "13619554740", "Оплата товаров по", "- 83.00 ₽", "- 83.00 ₽"],
            [None, None, "карте 4092 сумма 83.00"],
            [None, None, "в Mos.Transport"],
            [None, None, "MOSKVA RU дата 2026-"],
            [None, None, "09-21 время 18:16:16"],
            ["1"],
        ],
        [
            OZON_HEADER,
            # На последних страницах конвертер вставляет лишнюю колонку.
            ["20.09.2026 10:00:00", "4428324788", None, "Перевод клиенту Банка.",
             "+ 5 000.00 ₽", "+ 5 000.00 ₽"],
            [None, None, None, "Отправитель: Валентина. Без НДС."],
            ["Итого зачислений за период: 5 000.00 ₽"],
            ["Итого списаний за период: 83.00 ₽"],
            ["Исходящий остаток: 5 059.30 ₽"],
            ["С уважением,"],
            ["2"],
        ],
    ]


def test_reads_every_page_of_a_converted_statement():
    statement = get_connector("ozon").parse_statement(_paged_xlsx(_ozon_pages()), "vypiska.xlsx")

    # Вторая страница читается, хотя колонки в ней сдвинуты и шапки над
    # сдвинутой строкой нет.
    assert [op.amount for op in statement.operations] == [-83.0, 5000.0]
    assert statement.period_from == date(2026, 9, 20)


def test_wrapped_description_becomes_the_name_of_the_place():
    statement = get_connector("ozon").parse_statement(_paged_xlsx(_ozon_pages()), "vypiska.xlsx")

    payment, transfer = statement.operations
    assert payment.description == "Mos.Transport MOSKVA RU"
    assert classify(payment) == ("transport", "expense")
    # Перенос подхватывается и в сдвинутой строке, а подпись с последней
    # страницы к операции не приклеивается.
    assert transfer.description == "Перевод клиенту Банка. Отправитель: Валентина"


def test_balance_printed_in_the_statement_wins():
    statement = get_connector("ozon").parse_statement(_paged_xlsx(_ozon_pages()), "vypiska.xlsx")

    assert statement.closing_balance == 5059.30


def test_import_is_refused_when_it_misses_the_stated_totals():
    pages = _ozon_pages()
    pages[1][-5] = ["Итого зачислений за период: 500 000.00 ₽"]

    with pytest.raises(StatementParseError, match="не сходится с итогами"):
        get_connector("ozon").parse_statement(_paged_xlsx(pages), "vypiska.xlsx")


def test_refund_keeps_the_word_that_names_it():
    csv = (
        "Дата операции;Документ;Назначение платежа;Сумма операции\n"
        "21.09.2026;1;Возврат оплаты товаров по карте 3355 сумма 490.00 в "
        "DELIMOBIL MOSCOW RU дата 2026-09-21 время 10:00:00. Без НДС.;+ 490.00 ₽\n"
    )
    statement = get_connector("ozon").parse_statement(csv.encode(), "s.csv")
    operation = statement.operations[0]

    assert operation.description == "Возврат оплаты товаров: DELIMOBIL MOSCOW RU"
    assert classify(operation) == ("refunds", "income")


def test_parses_pdf_statement_best_effort():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    pdf.setFont("Helvetica", 10)
    y = 800
    for line in (
        "Vypiska po schetu Ozon Bank",
        "12.03.2026 Pyaterochka -1 234,56 RUB 48 765,44 RUB",
        "13.03.2026 Salary +180 000,00 RUB 228 765,44 RUB",
    ):
        pdf.drawString(40, y, line)
        y -= 20
    pdf.save()

    statement = get_connector("ozon").parse_statement(buffer.getvalue(), "statement.pdf")
    assert [op.amount for op in statement.operations] == [-1234.56, 180000.0]
    assert any("PDF" in warning for warning in statement.warnings)


def test_pdf_without_text_layer_gives_actionable_error():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    pdf.save()
    with pytest.raises(StatementParseError, match="CSV"):
        get_connector("ozon").parse_statement(buffer.getvalue(), "scan.pdf")


# --- categorisation ---------------------------------------------------------


@pytest.mark.parametrize(
    ("operation", "expected_category", "expected_type"),
    [
        (RawOperation(date(2026, 3, 1), -100, mcc="5411"), "groceries", "expense"),
        (RawOperation(date(2026, 3, 1), -100, mcc="5814"), "cafe", "expense"),
        (RawOperation(date(2026, 3, 1), -100, mcc="4121"), "transport", "expense"),
        (RawOperation(date(2026, 3, 1), -100, bank_category="Супермаркеты"), "groceries", "expense"),
        (RawOperation(date(2026, 3, 1), -100, merchant="Surf Coffee"), "cafe", "expense"),
        (RawOperation(date(2026, 3, 1), -100, merchant="DNS"), "gadgets", "expense"),
        (RawOperation(date(2026, 3, 1), 320, description="Кэшбэк за март"), "cashback", "income"),
        (RawOperation(date(2026, 3, 1), 180000, description="Зарплата"), "salary", "income"),
        (
            RawOperation(date(2026, 3, 1), -2000, description="Перевод между своими счетами"),
            "transfers",
            "transfer",
        ),
        (RawOperation(date(2026, 3, 1), -5000, description="Снятие наличных"), "cash", "expense"),
        # Эквайринг присылает название места латиницей — и это подавляющая
        # часть операций по карте.
        (
            RawOperation(date(2026, 3, 1), -66.82, description="MAGNIT MM TYUSHINO MOSCOW RU"),
            "groceries",
            "expense",
        ),
        (
            RawOperation(date(2026, 3, 1), -511, description="DELIMOBIL MOSCOW MOSCOW RU"),
            "transport",
            "expense",
        ),
        (
            RawOperation(date(2026, 3, 1), -83, description="Mos.Transport MOSKVA RU"),
            "transport",
            "expense",
        ),
        (
            RawOperation(date(2026, 3, 1), 52173, description="Заработная плата за Июль 2026 г."),
            "salary",
            "income",
        ),
        (
            RawOperation(date(2026, 3, 1), 214, description="Выплата кешбека по программе"),
            "cashback",
            "income",
        ),
        (
            RawOperation(date(2026, 3, 1), -1480, description="VV_8348_KCO_4 MOSCOW RU"),
            "groceries",
            "expense",
        ),
        (
            RawOperation(date(2026, 3, 1), -21580, description="EPGU MOSKVA RU", bank_category="Прочее"),
            "taxes",
            "expense",
        ),
        (RawOperation(date(2026, 3, 1), -100), "other", "expense"),
        (RawOperation(date(2026, 3, 1), 100), "other_income", "income"),
    ],
)
def test_classify(operation, expected_category, expected_type):
    assert classify(operation) == (expected_category, expected_type)


def test_wording_beats_mcc_for_cashback():
    # A cashback payout can carry the MCC of the shop that earned it.
    operation = RawOperation(date(2026, 3, 1), 320, description="Кэшбэк за март", mcc="5411")
    assert classify(operation) == ("cashback", "income")


# --- Open API client --------------------------------------------------------


def _transaction(transaction_id: str, amount: str, indicator: str) -> dict:
    return {
        "transactionId": transaction_id,
        "accountId": "200200",
        "Amount": {"amount": amount, "currency": "RUB"},
        "creditDebitIndicator": indicator,
        "status": "AcceptedSettlementCompleted",
        "bookingDateTime": "2026-03-12T15:15:13+00:00",
        "transactionInformation": "Покупка",
        "Creditor": {"name": "Пятерочка"},
        "MerchantInformation": {"merchantCategoryCode": "5411"},
    }


def test_open_banking_maps_transactions_to_operations():
    statement = statement_from_transactions(
        [_transaction("t1", "1234.56", "Debit"), _transaction("t2", "180000.00", "Credit")]
    )
    assert [op.amount for op in statement.operations] == [-1234.56, 180000.0]
    assert statement.operations[0].mcc == "5411"
    assert statement.operations[0].merchant == "Пятерочка"
    assert statement.operations[0].external_id == "t1"
    assert statement.period_from == date(2026, 3, 12)


@pytest.mark.asyncio
async def test_open_banking_client_follows_pagination():
    calls: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        calls.append(str(request.url))
        assert request.headers["Authorization"] == "Bearer token-123"
        assert "x-fapi-interaction-id" in request.headers
        if "page=2" in str(request.url):
            return httpx.Response(200, json={"Data": {"Transaction": [_transaction("t2", "10.00", "Credit")]}})
        return httpx.Response(
            200,
            json={
                "Data": {"Transaction": [_transaction("t1", "20.00", "Debit")]},
                "Links": {"next": "https://bank.example/accounts/200200/transactions?page=2"},
            },
        )

    client = OpenBankingClient(
        "https://bank.example",
        "token-123",
        transport=httpx.MockTransport(handler),
    )
    transactions = await client.fetch_transactions(
        "200200", since=date(2026, 3, 1), until=date(2026, 3, 31)
    )

    assert len(transactions) == 2
    assert len(calls) == 2
    assert "fromBookingDateTime=2026-03-01" in calls[0]


@pytest.mark.asyncio
async def test_open_banking_client_reports_auth_failures():
    client = OpenBankingClient(
        "https://bank.example",
        "expired",
        transport=httpx.MockTransport(lambda _r: httpx.Response(401, json={})),
    )
    with pytest.raises(OpenBankingError, match="401"):
        await client.list_accounts()


def test_open_banking_client_requires_base_url():
    with pytest.raises(OpenBankingError, match="выписки"):
        OpenBankingClient("", "token")


# --- ingest pipeline --------------------------------------------------------


async def _connection(db: AsyncSession) -> BankConnection:
    account = Account(name="Ozon Банк", account_type="bank", balance=0.0)
    db.add(account)
    await db.flush()
    connection = BankConnection(provider="ozon", label="Ozon Банк", account_id=account.id)
    db.add(connection)
    await db.flush()
    return connection


@pytest.mark.asyncio
async def test_ingest_creates_transactions_and_syncs_balance(db: AsyncSession):
    connection = await _connection(db)
    statement = get_connector("ozon").parse_statement(_csv_bytes(), "march.csv")

    run = await ingest_statement(db, connection, statement, source_name="march.csv")

    assert run.imported_count == 5
    assert run.duplicate_count == 0
    transactions = list((await db.execute(select(Transaction))).scalars())
    assert len(transactions) == 5
    assert {t.source for t in transactions} == {"import"}
    assert all("bank,ozon" == t.tags for t in transactions)

    account = await db.get(Account, connection.account_id)
    # The statement's own closing balance wins over summing the rows.
    assert account.balance == 226635.94
    assert connection.imported_total == 5
    assert connection.last_operation_on == date(2026, 3, 17)


@pytest.mark.asyncio
async def test_reimporting_the_same_statement_adds_nothing(db: AsyncSession):
    connection = await _connection(db)
    connector = get_connector("ozon")

    first = await ingest_statement(
        db, connection, connector.parse_statement(_csv_bytes(), "march.csv"), source_name="a.csv"
    )
    second = await ingest_statement(
        db, connection, connector.parse_statement(_csv_bytes(), "march.csv"), source_name="b.csv"
    )

    assert first.imported_count == 5
    assert second.imported_count == 0
    assert second.duplicate_count == 5
    assert len(list((await db.execute(select(Transaction))).scalars())) == 5


@pytest.mark.asyncio
async def test_genuine_same_day_duplicates_are_both_kept(db: AsyncSession):
    """Two identical coffees on one day are two operations, not a double import."""
    connection = await _connection(db)
    csv = (
        "Дата операции;Сумма операции;Описание операции\n"
        "12.03.2026;-200,00;Surf Coffee\n"
        "12.03.2026;-200,00;Surf Coffee\n"
    )
    connector = get_connector("ozon")

    first = await ingest_statement(
        db, connection, connector.parse_statement(csv.encode(), "a.csv"), source_name="a.csv"
    )
    assert first.imported_count == 2

    # Re-importing a statement that overlaps must still collapse to nothing new.
    second = await ingest_statement(
        db, connection, connector.parse_statement(csv.encode(), "b.csv"), source_name="b.csv"
    )
    assert second.imported_count == 0
    assert second.duplicate_count == 2


@pytest.mark.asyncio
async def test_overlapping_statement_imports_only_new_rows(db: AsyncSession):
    connection = await _connection(db)
    connector = get_connector("ozon")
    header = "Дата операции;Сумма операции;Описание операции\n"
    first_csv = header + "12.03.2026;-200,00;Surf Coffee\n"
    second_csv = first_csv + "13.03.2026;-300,00;Пятёрочка\n"

    await ingest_statement(
        db, connection, connector.parse_statement(first_csv.encode(), "a.csv"), source_name="a.csv"
    )
    run = await ingest_statement(
        db, connection, connector.parse_statement(second_csv.encode(), "b.csv"), source_name="b.csv"
    )

    assert (run.imported_count, run.duplicate_count) == (1, 1)
    operations = list((await db.execute(select(BankOperation))).scalars())
    assert len(operations) == 2


@pytest.mark.asyncio
async def test_external_ids_dedupe_across_reordered_statements(db: AsyncSession):
    connection = await _connection(db)
    statement = statement_from_transactions(
        [_transaction("t1", "100.00", "Debit"), _transaction("t2", "200.00", "Debit")]
    )
    reversed_statement = statement_from_transactions(
        [_transaction("t2", "200.00", "Debit"), _transaction("t1", "100.00", "Debit")]
    )

    await ingest_statement(db, connection, statement, source_name="api", source_kind="api")
    run = await ingest_statement(
        db, connection, reversed_statement, source_name="api", source_kind="api"
    )
    assert run.imported_count == 0
    assert run.duplicate_count == 2


@pytest.mark.asyncio
async def test_one_statement_in_two_formats_is_one_history(db: AsyncSession, monkeypatch):
    """Тот же год трат, выгруженный и таблицей, и PDF, — одна история.

    Номер документа виден по-разному: в таблице он лежит в своей колонке,
    в PDF короткий номер от описания не отличить. Пока личностью операции
    служил номер, тридцать девять зарплат завелись по второму разу.
    """
    connection = await _connection(db)
    connector = get_connector("ozon")
    table = (
        "Дата операции;Документ;Сумма операции;Описание операции\n"
        "10.08.2026;2079;+52 173,85;Заработная плата за Июль 2026 г.\n"
        "11.08.2026;12372063678;-83,00;Оплата проезда\n"
    )
    from_table = await ingest_statement(
        db, connection, connector.parse_statement(table.encode(), "god.csv"), source_name="a"
    )

    _pdf_text(monkeypatch, (
        "10.08.2026 09:00:00 2079 Заработная плата за Июль 2026 г.",
        "+ 52 173.85 ₽",
        "11.08.2026 09:00:00 12372063678 Оплата проезда",
        "- 83.00 ₽",
        "Итого зачислений за период: 52 173.85 ₽",
        "Итого списаний за период: 83.00 ₽",
    ))
    from_pdf = await ingest_statement(
        db, connection, connector.parse_statement(b"%PDF-1.4", "god.pdf"), source_name="b"
    )

    assert from_table.imported_count == 2
    assert (from_pdf.imported_count, from_pdf.duplicate_count) == (0, 2)
    assert len(list((await db.execute(select(Transaction))).scalars())) == 2


@pytest.mark.asyncio
async def test_history_written_by_the_old_rule_still_recognises_itself(db: AsyncSession):
    """Отпечатки, записанные когда личностью служил номер банка.

    Новое правило само по себе историю не чинит: пока в базе лежат старые
    отпечатки, та же выписка снова выглядит новой. Пересчёт идёт при запуске.
    """
    connection = await _connection(db)
    connector = get_connector("ozon")
    csv = (
        "Дата операции;Сумма операции;Описание операции\n"
        "12.03.2026;-200,00;Surf Coffee\n"
    )
    await ingest_statement(
        db, connection, connector.parse_statement(csv.encode(), "a.csv"), source_name="a"
    )
    stored = list((await db.execute(select(BankOperation))).scalars())
    for operation in stored:
        operation.fingerprint = "по-старому-от-номера-банка#0"
    await db.flush()

    assert await refresh_fingerprints(db) == 1
    assert await refresh_fingerprints(db) == 0, "пересчёт повторяться не должен"

    again = await ingest_statement(
        db, connection, connector.parse_statement(csv.encode(), "b.csv"), source_name="b"
    )
    assert (again.imported_count, again.duplicate_count) == (0, 1)


@pytest.mark.asyncio
async def test_pending_operations_are_held_back(db: AsyncSession):
    connection = await _connection(db)
    csv = (
        "Дата операции;Сумма операции;Описание операции;Статус\n"
        "12.03.2026;-200,00;Surf Coffee;В обработке\n"
        "12.03.2026;-300,00;Пятёрочка;Проведена\n"
    )
    statement = get_connector("ozon").parse_statement(csv.encode(), "a.csv")
    run = await ingest_statement(db, connection, statement, source_name="a.csv")

    assert (run.imported_count, run.skipped_count) == (1, 1)


# --- API --------------------------------------------------------------------


@pytest.mark.asyncio
async def test_providers_endpoint_lists_ozon(client: AsyncClient):
    response = await client.get("/api/v1/connections/providers")
    assert response.status_code == 200
    providers = response.json()
    ozon = next(p for p in providers if p["provider"] == "ozon")
    assert ozon["title"] == "Ozon Банк"
    assert ".csv" in ozon["statement_formats"]
    assert "Выписки и справки" in ozon["instructions"]


@pytest.mark.asyncio
async def test_connection_lifecycle_and_statement_upload(client: AsyncClient):
    created = await client.post("/api/v1/connections", json={"provider": "ozon"})
    assert created.status_code == 200
    connection = created.json()
    assert connection["account_id"] is not None
    assert connection["has_credentials"] is False

    upload = await client.post(
        f"/api/v1/connections/{connection['id']}/statement",
        files={"file": ("march.csv", _csv_bytes(), "text/csv")},
    )
    assert upload.status_code == 200
    run = upload.json()
    assert run["imported_count"] == 5
    assert run["status"] == "ok"

    transactions = (await client.get("/api/v1/transactions?limit=500")).json()
    imported = [t for t in transactions if t["source"] == "import"]
    assert len(imported) == 5
    assert {t["category"] for t in imported} >= {"groceries", "salary", "cafe", "cashback"}

    listed = (await client.get("/api/v1/connections")).json()
    assert listed[0]["imported_total"] == 5
    assert listed[0]["status"] == "idle"

    history = (await client.get(f"/api/v1/connections/{connection['id']}/imports")).json()
    assert history[0]["imported_count"] == 5


@pytest.mark.asyncio
async def test_broken_statement_returns_422_and_records_error(client: AsyncClient):
    connection = (await client.post("/api/v1/connections", json={"provider": "ozon"})).json()
    response = await client.post(
        f"/api/v1/connections/{connection['id']}/statement",
        files={"file": ("junk.csv", b"nothing useful here", "text/csv")},
    )
    assert response.status_code == 422
    assert "заголовка" in response.json()["detail"]

    listed = (await client.get("/api/v1/connections")).json()
    assert listed[0]["status"] == "error"
    assert listed[0]["last_error"]


@pytest.mark.asyncio
async def test_api_sync_without_credentials_fails_clearly(client: AsyncClient):
    connection = (
        await client.post("/api/v1/connections", json={"provider": "ozon", "mode": "api"})
    ).json()
    response = await client.post(f"/api/v1/connections/{connection['id']}/sync", json={})
    assert response.status_code in {400, 502}
    assert "выписки" in response.json()["detail"] or "токен" in response.json()["detail"]


@pytest.mark.asyncio
async def test_access_token_is_never_returned(client: AsyncClient):
    created = await client.post(
        "/api/v1/connections",
        json={"provider": "ozon", "mode": "api", "access_token": "super-secret"},
    )
    body = created.json()
    assert body["has_credentials"] is True
    assert "super-secret" not in created.text
    assert "access_token" not in body


# --- ways to quietly corrupt money -----------------------------------------


NEWEST_FIRST_CSV = """Дата операции;Сумма операции в валюте счёта;Валюта;Описание операции;Остаток после операции
31.03.2026;-100,00 ₽;RUB;Кофе;226 000,00
01.03.2026;-50,00 ₽;RUB;Булочка;10 000,00
"""


@pytest.mark.asyncio
async def test_two_cards_of_one_bank_keep_their_own_history(db: AsyncSession):
    """Одна и та же покупка бывает на двух картах: день, место и сумма совпали.

    Отпечаток жил в рамках провайдера, а колонка уникальна на всю базу, — и
    операции второй карты молча записывались в дубли первой.
    """
    first = await _connection(db)
    second = await _connection(db)
    statement = get_connector("ozon").parse_statement(_csv_bytes(), "march.csv")

    one = await ingest_statement(db, first, statement, source_name="march.csv")
    two = await ingest_statement(
        db,
        second,
        get_connector("ozon").parse_statement(_csv_bytes(), "march.csv"),
        source_name="march.csv",
    )

    assert one.imported_count == 5
    assert two.imported_count == 5, "вторая карта потеряла свои операции"
    assert two.duplicate_count == 0
    operations = list((await db.execute(select(BankOperation))).scalars())
    assert len({op.connection_id for op in operations}) == 2


def test_the_closing_balance_comes_from_the_latest_operation():
    """Ozon выгружает свежее сверху, и нижняя строка файла — самая старая."""
    statement = rows_to_statement(read_csv_rows(NEWEST_FIRST_CSV.encode()))

    assert statement.closing_balance == 226000.0


@pytest.mark.asyncio
async def test_an_older_statement_does_not_roll_the_balance_back(db: AsyncSession):
    connection = await _connection(db)
    connector = get_connector("ozon")
    await ingest_statement(
        db, connection, connector.parse_statement(_csv_bytes(), "march.csv"),
        source_name="march.csv",
    )
    account = await db.get(Account, connection.account_id)
    fresh = account.balance

    old = """Дата операции;Сумма операции в валюте счёта;Валюта;Описание операции;Остаток после операции
05.01.2026;-700,00 ₽;RUB;Аптека;9 300,00
"""
    run = await ingest_statement(
        db, connection, connector.parse_statement(old.encode(), "january.csv"),
        source_name="january.csv",
    )

    assert run.imported_count == 1, "старые операции всё равно нужно записать"
    assert account.balance == fresh, "остаток откатился в январь"
    assert "остаток" in run.warnings_json


def test_a_ruble_purchase_partly_paid_with_points_survives():
    """«Оплачено 500 баллами и 1200 ₽» — это настоящие 1200 рублей."""
    csv = """Дата операции;Сумма операции в валюте счёта;Валюта;Категория;Описание операции
10.03.2026;-1 200,00 ₽;RUB;Покупки;Ozon заказ, оплачено 500 баллами и 1200 ₽
11.03.2026;+300,00;БАЛЛЫ;Баллы;Начислен кэшбэк
12.03.2026;-90,00 ₽;RUB;Бонусная программа;Плата за подписку Premium
"""
    statement = get_connector("ozon").parse_statement(csv.encode(), "march.csv")

    kept = {op.description for op in statement.operations}
    assert len(statement.operations) == 2, kept
    assert any("оплачено 500 баллами" in d for d in kept)
    assert any("Premium" in d for d in kept)

def _pdf_text(monkeypatch, lines: tuple[str, ...]) -> None:
    """Подменить только текстовый слой: шрифты reportlab не знают кириллицу."""
    from app.connectors import pdf_statement

    monkeypatch.setattr(pdf_statement, "extract_pdf_text", lambda data: "\n".join(lines))


def test_pdf_totals_do_not_become_operations(monkeypatch):
    """Итоги и остатки начинаются с даты, но деньгами не являются.

    Записанные как операции, они удваивают траты: сначала в истории, потом в
    остатке.
    """
    _pdf_text(monkeypatch, (
        "Выписка по счёту Ozon Банк",
        "01.03.2026 - 31.03.2026 Поступления 50 000,00 ₽",
        "01.03.2026 Входящий остаток 12 345,00 ₽",
        "12.03.2026 Пятёрочка -1 234,56 ₽",
        "31.03.2026 Итого списаний 1 234,56 ₽",
        "31.03.2026 Исходящий остаток 11 110,44 ₽",
    ))

    statement = get_connector("ozon").parse_statement(b"%PDF-1.4", "statement.pdf")

    assert [op.amount for op in statement.operations] == [-1234.56]
    assert any("итогами" in w for w in statement.warnings)


def test_pdf_reads_an_operation_whose_amount_is_printed_under_it(monkeypatch):
    """Так выписку печатает Ozon: перенос описания, а сумма — отдельной строкой.

    Пока сумма требовалась в строке с датой, из годовой выписки читалась
    горстка операций, а остальные молча пропадали.
    """
    _pdf_text(monkeypatch, (
        "Дата операции Документ Назначение платежа",
        "21.09.2026 18:51:01 13619554740 Оплата товаров по ",
        "карте 4092 сумма 83.00 ",
        "в Mos.Transport ",
        "MOSKVA RU дата 2026-",
        "09-21 время 18:16:16",
        "- 83.00 ₽ - 83.00 ₽",
        "Итого списаний за период: 83.00 ₽",
        "Исходящий остаток: 59.30 ₽",
    ))

    statement = get_connector("ozon").parse_statement(b"%PDF-1.4", "vypiska.pdf")
    operation = statement.operations[0]

    assert operation.amount == -83.0
    assert operation.description == "Mos.Transport MOSKVA RU"
    assert operation.external_id == "13619554740"
    assert statement.closing_balance == 59.30


def test_pdf_document_number_is_not_an_amount(monkeypatch):
    """Строка с датой кончается номером документа, а описание — со следующей.

    Принятый за сумму, одиннадцатизначный номер заводил операцию на миллиард
    рублей — и таких в годовой выписке были сотни.
    """
    _pdf_text(monkeypatch, (
        "17.09.2026 09:52:58 1349425902",
        "Перевод клиенту Банка",
        "287",
        "Дата операции Документ Назначение платежа",
        "- 1 000.00 ₽ - 1 000.00 ₽",
        "Итого списаний за период: 1 000.00 ₽",
    ))

    statement = get_connector("ozon").parse_statement(b"%PDF-1.4", "vypiska.pdf")

    assert [op.amount for op in statement.operations] == [-1000.0]
    # Разрыв страницы посреди операции в описание не попадает.
    assert statement.operations[0].description == "Перевод клиенту Банка"


def test_pdf_short_document_number_is_not_an_identifier(monkeypatch):
    """Номер документа уходит из описания, но удостоверяет не всякий.

    В таблице этот номер лежит в своей колонке, поэтому в описании его быть
    не должно — иначе одна и та же выписка в двух форматах выглядит двумя
    разными историями. Но короткий номер в выписке повторяется, и сверять по
    нему операции нельзя: разные операции слились бы в одну.
    """
    _pdf_text(monkeypatch, (
        "10.08.2026 09:00:00 2079 Для зачисления на счет",
        "Заработная плата за Июль 2026 г.",
        "+ 52 173.85 ₽",
        "Итого зачислений за период: 52 173.85 ₽",
    ))

    operation = get_connector("ozon").parse_statement(b"%PDF-1.4", "vypiska.pdf").operations[0]

    assert operation.description == "Для зачисления на счет Заработная плата за Июль 2026 г."
    assert operation.external_id == ""
    assert classify(operation) == ("salary", "income")


def test_pdf_statement_period_is_not_taken_from_an_operation(monkeypatch):
    """«За период с … по …» пишут и в описании операции.

    Годовая выписка так объявила себя месячной: период взялся из строки про
    выплату кешбэка.
    """
    _pdf_text(monkeypatch, (
        "10.01.2026 09:00:00 Выплата кешбека за период с 10.12.2025 по 10.01.2026",
        "+ 214.00 ₽",
        "23.09.2025 09:00:00 Оплата товаров по карте 4092 сумма 83.00 в MAGNIT RU",
        "- 83.00 ₽",
    ))

    statement = get_connector("ozon").parse_statement(b"%PDF-1.4", "vypiska.pdf")

    # Период накрывает обе операции, а не одну декабрьскую строку.
    assert statement.period_from == date(2025, 9, 23)
    assert statement.period_to == date(2026, 1, 10)


def test_pdf_income_without_a_sign_is_not_turned_into_a_loss(monkeypatch):
    """Без знака направление читается из формулировки, а не назначается."""
    _pdf_text(monkeypatch, (
        "Выписка",
        "13.03.2026 Зачисление зарплаты 180 000,00 ₽",
        "14.03.2026 Оплата Surf Coffee 450,00 ₽",
    ))

    statement = get_connector("ozon").parse_statement(b"%PDF-1.4", "statement.pdf")

    assert [op.amount for op in statement.operations] == [180000.0, -450.0]


def test_a_date_inside_a_description_does_not_start_an_operation(monkeypatch):
    """«…по обращению от 16.07.2026. Сумма 19779-» — это перенос, а не операция.

    Приняв такую строку за начало новой операции, разбор заводил зарплату
    задним числом, а настоящую, чьё описание он оборвал, терял. Денег в сумме
    столько же, поэтому итоги сходились и подмена ничем себя не выдавала.
    """
    _pdf_text(monkeypatch, (
        "06.08.2026 16:37:46 144014 Для зачисления на счет",
        "Кириллов Кирилл Владимирович. Компенсация расходов по",
        "обращению от",
        "16.07.2026. Сумма 19779-",
        "00 В т.ч. Без налога (НДС)",
        "+ 19 779.00 ₽ + 19 779.00 ₽",
        "Итого зачислений за период: 19 779.00 ₽",
    ))

    statement = get_connector("ozon").parse_statement(b"%PDF-1.4", "vypiska.pdf")

    assert [(op.occurred_on, op.amount) for op in statement.operations] == [
        (date(2026, 8, 6), 19779.0)
    ]


def test_pdf_document_number_goes_away_even_before_a_numeric_purpose(monkeypatch):
    """Назначение платежа тоже начинается с цифр — номер узнаётся по месту.

    После даты и времени первое число — номер документа, каким бы ни было
    описание. В таблице этот номер лежит в своей колонке, и пока он оставался
    в тексте, зарплата из PDF и зарплата из XLSX были разными операциями.
    """
    _pdf_text(monkeypatch, (
        "08.09.2026 09:00:00 2506 187/№187/249/2/ОП/ 26/84/Для зачисления",
        "на счет Кириллова К.В. Заработная плата август",
        "+ 37 986.12 ₽",
        "Итого зачислений за период: 37 986.12 ₽",
    ))

    operation = get_connector("ozon").parse_statement(b"%PDF-1.4", "v.pdf").operations[0]

    assert operation.description.startswith("187/№187/249/2/ОП/")
    assert operation.external_id == ""


def test_a_pdf_that_hides_the_direction_is_refused(monkeypatch):
    """Две колонки банк напечатал, а в тексте от них ничего не осталось."""
    _pdf_text(monkeypatch, (
        "Выписка",
        "12.03.2026 ТОРГОВАЯ ТОЧКА 1 234,56 ₽",
        "13.03.2026 ООО РОМАШКА 180 000,00 ₽",
    ))

    with pytest.raises(StatementParseError, match="CSV"):
        get_connector("ozon").parse_statement(b"%PDF-1.4", "statement.pdf")


# --- адрес банка как окно внутрь сети --------------------------------------


@pytest.mark.parametrize(
    "url",
    [
        "http://127.0.0.1:8820/v1",
        "http://localhost:8800/api",
        "http://10.0.0.5/api",
        "http://[::1]/api",
        "http://169.254.169.254/latest/meta-data",
        "file:///etc/passwd",
        "https://",
    ],
)
def test_the_bank_address_cannot_point_inside_our_network(url):
    """Клиент возвращает тело того, до кого дотянулся.

    Без проверки поле «адрес Открытого API» — это окно на всё, что слушает на
    машине: ассистент, мост телефона с историей здоровья, служба метаданных.
    """
    with pytest.raises(OpenBankingError):
        OpenBankingClient(url, "token")


def test_a_real_bank_address_is_accepted():
    client = OpenBankingClient("https://bank.example/open-api/", "token")

    assert client is not None


@pytest.mark.asyncio
async def test_the_access_token_does_not_follow_the_bank_elsewhere():
    """Links.next вёл куда угодно, а токен уходил с каждым запросом."""
    seen: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(str(request.url))
        return httpx.Response(
            200,
            json={
                "Data": {"Transaction": []},
                "Links": {"next": "https://attacker.example/collect"},
            },
        )

    client = OpenBankingClient(
        "https://bank.example", "token-123", transport=httpx.MockTransport(handler)
    )

    with pytest.raises(OpenBankingError, match="токен"):
        await client.list_accounts()
    assert all("attacker" not in url for url in seen)


@pytest.mark.asyncio
async def test_a_relative_next_page_still_works():
    pages: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        pages.append(str(request.url))
        if "page=2" in str(request.url):
            return httpx.Response(200, json={"Data": {"Account": [{"accountId": "2"}]}})
        return httpx.Response(
            200,
            json={
                "Data": {"Account": [{"accountId": "1"}]},
                "Links": {"next": "accounts?page=2"},
            },
        )

    client = OpenBankingClient(
        "https://bank.example", "token", transport=httpx.MockTransport(handler)
    )

    assert len(await client.list_accounts()) == 2
    assert pages[1] == "https://bank.example/accounts?page=2"


def test_a_table_with_absurdly_many_rows_is_refused():
    """Размер файла ничего не говорит о таблице внутри."""
    from app.connectors.tabular import MAX_ROWS

    huge = "Дата операции;Сумма операции в валюте счёта;Описание операции\n" + (
        "12.03.2026;-1,00;Кофе\n" * (MAX_ROWS + 5)
    )

    with pytest.raises(StatementParseError, match="строк"):
        read_csv_rows(huge.encode())


def _cyrillic_font() -> str | None:
    """Шрифт с кириллицей: встроенные в reportlab её не содержат.

    Без него русский текст уходит в PDF пустыми глифами, и проверять словарь
    разбора — русский целиком — было бы не на чем.
    """
    from reportlab.pdfbase import pdfmetrics, ttfonts

    for path in (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/TTF/DejaVuSans.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
    ):
        if Path(path).exists():
            try:
                pdfmetrics.registerFont(ttfonts.TTFont("cyr", path))
                return "cyr"
            except Exception:  # noqa: BLE001 - шрифт есть, но не читается
                continue
    return None


def _pdf(lines: list[str], font: str | None = None) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    face = font or "Helvetica"
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    pdf.setFont(face, 9)
    y = 800
    for line in lines:
        pdf.drawString(40, y, line)
        y -= 12
        if y < 40:
            pdf.showPage()
            pdf.setFont(face, 9)
            y = 800
    pdf.save()
    return buffer.getvalue()


def test_the_currency_sign_is_printed_once_in_the_header_not_on_every_row():
    """Пока знак валюты был обязателен, выписка теряла почти все операции."""
    statement = get_connector("ozon").parse_statement(
        _pdf([
            "Data Opisanie Summa, RUB Ostatok",
            "12.03.2026 Pyaterochka -1 234,56 48 765,44",
            "13.03.2026 Zachislenie zarplaty +180 000,00 228 765,44",
            "14.03.2026 Oplata Yandex Taxi -450,00 228 315,44",
        ]),
        "statement.pdf",
    )

    assert [op.amount for op in statement.operations] == [-1234.56, 180000.0, -450.0]


def test_a_page_number_is_not_an_operation():
    statement = get_connector("ozon").parse_statement(
        _pdf([
            "12.03.2026 Pyaterochka -1 234,56",
            "13.03.2026 Salary +180 000,00",
            "01.03.2026 Stranica 2",
        ]),
        "statement.pdf",
    )

    assert [op.amount for op in statement.operations] == [-1234.56, 180000.0]


def test_a_layout_it_did_not_understand_is_refused_not_guessed():
    """Полгода трат не могут превратиться в две операции незаметно."""
    lines = [f"{day:02d}.03.2026 Pokupka v magazine bez summy" for day in range(1, 29)]
    lines += ["12.03.2026 Pyaterochka -1 234,56 RUB"]

    with pytest.raises(StatementParseError, match="не понял вёрстку"):
        get_connector("ozon").parse_statement(_pdf(lines), "statement.pdf")


def test_a_credit_is_not_a_debit_just_because_the_bank_omits_the_plus():
    """Банк ставит минус на списания; строка без знака — это приход.

    Иначе зарплата в такой выписке ошибается на две своих суммы.
    """
    statement = get_connector("ozon").parse_statement(
        _pdf([
            "Data Opisanie Summa, RUB Ostatok",
            "02.03.2026 Pyaterochka -1 234,56 48 765,44",
            "05.03.2026 Zachislenie zarplaty 180 000,00 228 765,44",
            "07.03.2026 Ozon -3 299,00 225 466,44",
        ]),
        "statement.pdf",
    )

    assert [op.amount for op in statement.operations] == [-1234.56, 180000.0, -3299.0]


def test_without_any_signs_the_wording_still_decides():
    font = _cyrillic_font()
    if font is None:
        pytest.skip("нет шрифта с кириллицей — словарь разбора проверить нечем")

    statement = get_connector("ozon").parse_statement(
        _pdf([
            "02.03.2026 Оплата Пятёрочка 1 234,56",
            "05.03.2026 Зачисление зарплаты 180 000,00",
        ], font),
        "statement.pdf",
    )

    assert [op.amount for op in statement.operations] == [-1234.56, 180000.0]
