"""Tolerant reader for tabular bank statements (CSV / XLSX).

Banks rename statement columns freely and ship them in CP1251, UTF-8 with a BOM
or UTF-16, with `;` or `,` as a delimiter and a few decorative rows above the
header. Rather than hardcoding one layout, this module scores every header cell
against a pattern table and picks the best column per field. That way a small
change in the bank's export keeps working.
"""

from __future__ import annotations

import codecs
import csv
import io
import re
from collections import Counter
from datetime import date
from typing import Any

from app.connectors.base import (
    ParsedStatement,
    RawOperation,
    StatementParseError,
    check_against_stated,
    parse_money,
    parse_statement_date,
    stated_totals,
    unique_reference,
)

# field -> ((header fragment, score), ...). Highest score wins the column.
FIELD_PATTERNS: dict[str, tuple[tuple[str, int], ...]] = {
    "occurred_on": (
        ("датаоперации", 100), ("датасовершения", 96), ("датаивремяоперации", 95),
        ("датапокупки", 92), ("дататранзакции", 90), ("датаивремя", 80),
        ("датаплатежа", 70), ("датадокумента", 60), ("датаобработки", 40),
        ("датасписания", 40), ("датапроведения", 40), ("дата", 30),
        ("date", 30), ("datetime", 25),
    ),
    "amount": (
        ("суммаоперацииввалютесчета", 100), ("суммаввалютесчета", 98),
        ("суммавалютесчета", 96), ("суммаврублях", 94), ("суммавруб", 92),
        ("суммаоперации", 85), ("суммаплатежа", 80), ("суммасучетомкомиссии", 78),
        ("сумма", 60), ("amount", 60),
    ),
    "credit": (
        ("приход", 100), ("зачисление", 96), ("поступление", 94), ("пополнение", 90),
        ("кредит", 70), ("credit", 70),
    ),
    "debit": (
        ("расход", 100), ("списание", 96), ("списано", 94), ("дебет", 70),
        ("debit", 70), ("выдача", 60),
    ),
    "direction": (
        ("типоперации", 100), ("направление", 92), ("видоперации", 85),
        ("приходрасход", 80), ("тип", 50),
    ),
    "description": (
        ("описаниеоперации", 100), ("назначениеплатежа", 96), ("описание", 94),
        ("назначение", 85), ("комментарий", 80), ("детали", 70),
        ("наименованиеоперации", 85), ("description", 90),
    ),
    "merchant": (
        ("контрагент", 100), ("торговаяточка", 96), ("мерчант", 95), ("merchant", 95),
        ("магазин", 90), ("местосовершения", 88), ("продавец", 85), ("получатель", 70),
        ("плательщик", 60),
    ),
    "bank_category": (("категорияоперации", 100), ("категория", 96), ("category", 90)),
    "mcc": (("mccкод", 100), ("mcc", 96), ("мсс", 90), ("кодкатегории", 80)),
    "currency": (
        ("валютасчета", 100), ("валютаоперации", 90), ("валюта", 85), ("currency", 85),
    ),
    "status": (("статусоперации", 100), ("статус", 96), ("status", 90)),
    "external_id": (
        ("номероперации", 100), ("идентификатороперации", 98), ("идентификатор", 90),
        ("референс", 85), ("reference", 85), ("номердокумента", 70), ("authcode", 60),
        ("id", 55), ("документ", 50),
    ),
    "balance": (
        ("остатокпослеоперации", 100), ("баланспослеоперации", 98), ("остаток", 90),
        ("balance", 85),
    ),
}

DECLINED_MARKERS = ("отклон", "отказ", "неуспеш", "не выполн", "cancel", "fail", "reject", "возврат платежа")
PENDING_MARKERS = ("обработ", "ожид", "hold", "pending", "авториз", "не подтвержд")
DEBIT_MARKERS = ("списание", "расход", "оплата", "покупка", "debit", "перевод со счета")
CREDIT_MARKERS = ("поступление", "приход", "зачисление", "пополнение", "credit", "внесение")

_LEGACY_ENCODINGS = ("cp1251", "cp866", "koi8-r")
_MAX_HEADER_SCAN = 30


def normalize_header(value: Any) -> str:
    text = str(value or "").lower().replace("ё", "е")
    return "".join(ch for ch in text if ch.isalnum())


def _score_column(field: str, header: str) -> int:
    best = 0
    for fragment, score in FIELD_PATTERNS[field]:
        if len(fragment) <= 3:
            if header == fragment:
                best = max(best, score)
        elif fragment in header:
            best = max(best, score)
    return best


def map_header(row: list[Any]) -> tuple[dict[str, int], int]:
    """Map fields onto column indexes by one header row, with its total score."""
    headers = [normalize_header(cell) for cell in row]
    if not any(headers):
        return {}, 0

    # Per field, remember the best-scoring column; per column, remember the
    # best-scoring field. A column may only serve one field.
    per_field: dict[str, tuple[int, int]] = {}
    for col_index, header in enumerate(headers):
        if not header:
            continue
        for field in FIELD_PATTERNS:
            score = _score_column(field, header)
            if score and score > per_field.get(field, (0, -1))[0]:
                per_field[field] = (score, col_index)

    claimed: dict[int, tuple[str, int]] = {}
    for field, (score, col_index) in per_field.items():
        current = claimed.get(col_index)
        if current is None or score > current[1]:
            claimed[col_index] = (field, score)

    mapping = {field: col_index for col_index, (field, _) in claimed.items()}
    if "occurred_on" not in mapping or not {"amount", "credit", "debit"} & mapping.keys():
        return {}, 0
    return mapping, sum(score for _, score in claimed.values())


def detect_columns(rows: list[list[Any]]) -> tuple[int, dict[str, int]]:
    """Find the header row and map fields onto column indexes."""
    best_row, best_mapping, best_total = -1, {}, 0

    for row_index, row in enumerate(rows[:_MAX_HEADER_SCAN]):
        mapping, total = map_header(row)
        if mapping and total > best_total:
            best_row, best_mapping, best_total = row_index, mapping, total

    if best_row < 0:
        guessed = guess_columns(rows)
        if guessed:
            return -1, guessed
        raise StatementParseError(_nothing_to_read(rows))
    return best_row, best_mapping


# Строки, похожие на операции, ещё не выписка: пара дат с суммами найдётся и в
# справке, и в счёте. Год трат — это сотни строк, поэтому колонки угадываются
# только когда одна и та же пара колонок повторяется много раз подряд.
MIN_GUESSED_ROWS = 10
_MONEY_MARKS = re.compile(r"[.,]\d{1,2}(?!\d)|[₽$€]|руб|[-+\u2212]\s?\d|\d[\s\u00a0]\d{3}", re.I)


def _looks_like_money(cell: Any) -> bool:
    """Сумма, а не номер документа и не количество.

    Номер документа парсится как число ничуть не хуже суммы, поэтому одного
    разбора мало: деньги себя выдают копейками, знаком, разрядами или валютой.
    """
    if isinstance(cell, bool) or cell is None:
        return False
    if isinstance(cell, float):
        return cell != 0
    if isinstance(cell, int):
        return False
    text = str(cell).strip()
    return bool(text) and parse_money(text) not in (None, 0) and bool(_MONEY_MARKS.search(text))


def guess_columns(rows: list[list[Any]]) -> dict[str, int]:
    """Колонки по содержимому строк, когда шапки в файле нет.

    Банки выгружают и такое: «справка о движении средств», переведённая в
    таблицу, теряет шапку целиком. Раньше разбор отвечал «проверьте, что это
    выписка» — про файл, который выпиской и был.
    """
    votes: Counter[tuple[int, int]] = Counter()
    for row in rows:
        date_at = next((i for i, cell in enumerate(row) if parse_statement_date(cell)), None)
        if date_at is None:
            continue
        money_at = next(
            (i for i, cell in enumerate(row) if i != date_at and _looks_like_money(cell)), None
        )
        if money_at is not None:
            votes[(date_at, money_at)] += 1

    if not votes:
        return {}
    (date_at, money_at), seen = votes.most_common(1)[0]
    if seen < MIN_GUESSED_ROWS:
        return {}

    mapping = {"occurred_on": date_at, "amount": money_at}
    described = _wordiest_column(rows, skip={date_at, money_at})
    if described is not None:
        mapping["description"] = described
    return mapping


def _wordiest_column(rows: list[list[Any]], skip: set[int]) -> int | None:
    """Колонка с самым длинным текстом — в выписке это назначение платежа."""
    letters: Counter[int] = Counter()
    for row in rows:
        for index, cell in enumerate(row):
            if index in skip or cell is None:
                continue
            letters[index] += sum(1 for ch in str(cell) if ch.isalpha())
    return max(letters, key=letters.get) if letters else None


def _nothing_to_read(rows: list[list[Any]]) -> str:
    """Отказ, по которому видно, что за файл прочитали."""
    sample = ""
    for row in rows[:_MAX_HEADER_SCAN]:
        text = " | ".join(str(cell).strip() for cell in row if str(cell or "").strip())
        if len(text) > len(sample):
            sample = text
    seen = f"Строк в файле: {len(rows)}."
    if sample:
        seen += f" Самая содержательная: «{sample[:160]}»."
    return (
        "Не нашёл в файле ни строки заголовка с датой и суммой операции, ни строк, "
        f"похожих на операции. {seen} Похоже, это справка, счёт или квитанция, "
        "а не выписка по счёту."
    )


# A year of card history is a few thousand rows. These caps exist because the
# file size says nothing about the table inside: a 5 MB XLSX is a zip, and it
# can declare millions of rows that cost gigabytes once expanded row by row.
MAX_ROWS = 100_000
MAX_COLUMNS = 200


def _limited(rows: Any, already: int = 0) -> list[list[Any]]:
    result: list[list[Any]] = []
    for row in rows:
        if already + len(result) >= MAX_ROWS:
            raise StatementParseError(
                f"В выписке больше {MAX_ROWS} строк — это не похоже на выписку "
                "по счёту. Выгрузите период поменьше."
            )
        result.append(list(row)[:MAX_COLUMNS])
    return result


def read_csv_rows(data: bytes) -> list[list[str]]:
    text = _decode(data)
    sample = "\n".join(text.splitlines()[:20])
    delimiter = max(";,\t", key=sample.count) if sample else ","
    if sample.count(delimiter) == 0:
        delimiter = ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return _limited(reader)


def read_xlsx_rows(data: bytes) -> list[list[Any]]:
    """Все листы книги подряд, одной таблицей.

    Один лист — это предположение о том, что банк выгружает таблицу. Ozon
    выдаёт «справку о движении средств» как PDF, и переведённая в XLSX она
    приходит постранично: 288 листов по странице в каждом, с повторённой
    шапкой. Год трат лежал на втором листе и дальше, а читался только первый.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - openpyxl is a hard dependency
        raise StatementParseError("Для чтения XLSX нужен пакет openpyxl") from exc

    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        rows: list[list[Any]] = []
        for name in workbook.sheetnames:
            sheet = workbook[name]
            # Размер листа в файле — это заявление, а не факт. Конвертеры
            # пишут «A1:A1», и режим чтения верит: из выписки на тысячи строк
            # приходила одна, и разбор жаловался на отсутствие шапки.
            sheet.reset_dimensions()
            rows.extend(_limited(sheet.iter_rows(values_only=True), len(rows)))
        return rows
    finally:
        workbook.close()


def _decode(data: bytes) -> str:
    """Decode statement bytes, honouring BOMs and falling back to CP1251.

    UTF-8 is self-validating, so a successful strict decode is trustworthy. The
    check has to come first: decoding UTF-8 Cyrillic as CP1251 always succeeds
    and yields twice as many "Cyrillic" characters, so any heuristic that just
    counts letters picks the mojibake.
    """
    if data.startswith(codecs.BOM_UTF8):
        return data.decode("utf-8-sig")
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        pass
    for encoding in _LEGACY_ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _cell(row: list[Any], mapping: dict[str, int], field: str) -> Any:
    index = mapping.get(field)
    if index is None or index >= len(row):
        return None
    return row[index]


def _text(row: list[Any], mapping: dict[str, int], field: str) -> str:
    value = _cell(row, mapping, field)
    return "" if value is None else str(value).strip()


def _closing_balance(
    balances: list[tuple[date, int, float]], dates: list[date]
) -> float | None:
    """The balance after the latest operation, not after the last row.

    Ozon and Tinkoff export newest first, so the bottom row of the file is the
    oldest operation — and its balance is weeks stale. Reading it as the closing
    balance overwrote the account with a number from the past.
    """
    if not balances:
        return None
    latest = max(day for day, _, _ in balances)
    newest_first = dates[0] > dates[-1]
    same_day = sorted(
        ((index, value) for day, index, value in balances if day == latest),
        reverse=not newest_first,
    )
    return same_day[0][1]


def rows_to_statement(rows: list[list[Any]]) -> ParsedStatement:
    """Turn spreadsheet rows into a `ParsedStatement`."""
    header_index, mapping = detect_columns(rows)
    guessed = header_index < 0
    header = [] if guessed else [normalize_header(cell) for cell in rows[header_index]]
    statement = ParsedStatement()
    unreadable = 0
    balances: list[tuple[date, int, float]] = []
    last: RawOperation | None = None
    # Колонка, из которой пришло описание последней операции: перенос строки
    # принимается только из неё. Иначе к последней операции приклеивается
    # подпись «С уважением,» с последней страницы.
    tail = mapping.get("description", -1)

    for row in rows[header_index + 1 :]:
        if not any(str(cell).strip() for cell in row if cell is not None):
            continue
        if _repeats_header(row, header):
            # Постраничная выписка повторяет шапку на каждой странице, и на
            # странице колонки могут стоять иначе: на последней странице этой
            # выписки они сдвинуты на одну, и её операции читались как пустые.
            again, _ = map_header(row)
            mapping = again or mapping
            continue

        occurred_on = parse_statement_date(_cell(row, mapping, "occurred_on"))
        columns = mapping
        amount = _row_amount(row, columns)
        if occurred_on is not None and not amount:
            columns = _realign(row, mapping)
            amount = _row_amount(row, columns)

        if occurred_on is None or amount is None or amount == 0:
            if _continues(row, tail, last):
                _extend(last, str(row[tail]).strip())
            else:
                unreadable += 1
            continue

        status = _text(row, columns, "status").lower()
        if any(marker in status for marker in DECLINED_MARKERS):
            continue

        description = _text(row, columns, "description")
        merchant = _text(row, columns, "merchant") or description[:200]
        currency = (_text(row, columns, "currency") or "RUB").upper()[:8] or "RUB"
        mcc = "".join(ch for ch in _text(row, columns, "mcc") if ch.isdigit())

        last = RawOperation(
            occurred_on=occurred_on,
            amount=amount,
            description=description or merchant,
            merchant=merchant,
            currency=currency,
            external_id=unique_reference(_text(row, columns, "external_id")),
            mcc=mcc,
            bank_category=_text(row, columns, "bank_category"),
            is_pending=any(marker in status for marker in PENDING_MARKERS),
            raw={"row": [str(cell) if cell is not None else "" for cell in row]},
        )
        statement.operations.append(last)
        tail = columns.get("description", -1)
        balance = parse_money(_cell(row, columns, "balance"))
        if balance is not None:
            balances.append((occurred_on, len(statement.operations) - 1, balance))

    if not statement.operations:
        raise StatementParseError(
            "Заголовок выписки распознан, но ни одной операции прочитать не удалось."
        )

    dates = [op.occurred_on for op in statement.operations]
    statement.period_from = min(dates)
    statement.period_to = max(dates)
    stated = stated_totals(cell for row in rows for cell in row)
    # Напечатанный в выписке остаток вернее посчитанного по строкам.
    statement.closing_balance = stated.get("closing", _closing_balance(balances, dates))
    if guessed:
        statement.warnings.append(
            "Шапки в файле нет — колонки определены по содержимому строк. "
            "Сверьте несколько операций на странице «Операции»."
        )
    if unreadable:
        statement.warnings.append(f"Пропущено нечитаемых строк: {unreadable}")
    check_against_stated(statement, stated)
    return statement


# Назначение платежа в постраничной выписке не помещается в ячейку и
# переносится на следующие строки. В первой ячейке остаётся «Оплата товаров
# по», а магазин — во второй и третьей: без склейки список операций
# превращается в тысячу одинаковых строк, и категорию определить не по чему.
DESCRIPTION_LIMIT = 400


def _repeats_header(row: list[Any], header: list[str]) -> bool:
    cells = [normalize_header(cell) for cell in row]
    filled = [cell for cell in cells if cell]
    return bool(filled) and all(cell in header for cell in filled)


def _continues(row: list[Any], column: int, last: RawOperation | None) -> bool:
    """Продолжение описания: ни даты, ни суммы, только текст в той же колонке.

    В остальных ячейках иногда оказывается номер страницы: постраничный
    конвертер кладёт его в соседнюю колонку. Номер документа с ним не
    спутать — он одиннадцатизначный.
    """
    if last is None or column < 0 or column >= len(row):
        return False
    if not str(row[column] or "").strip():
        return False
    for index, cell in enumerate(row):
        if index == column or cell is None:
            continue
        text = str(cell).strip()
        if text and not (text.isdigit() and len(text) <= 4):
            return False
    return True


def _realign(row: list[Any], mapping: dict[str, int]) -> dict[str, int]:
    """Колонки этой строки, если они сдвинуты относительно шапки.

    Постраничный конвертер вставляет на последних страницах лишнюю пустую
    колонку, причём посреди страницы и без новой шапки. Такие строки читались
    как пустые: шесть операций пропадали, и обороты не сходились с итогами,
    которые банк напечатал в самой выписке.
    """
    if "amount" not in mapping:
        return mapping
    money = next(
        (index for index in range(mapping["amount"], len(row)) if parse_money(row[index])),
        -1,
    )
    if money < 0 or money == mapping["amount"]:
        return mapping

    shifted = dict(mapping, amount=money)
    said = mapping.get("description")
    if said is not None and not str(_cell(row, mapping, "description") or "").strip():
        text = next(
            (index for index in range(said, money)
             if str(row[index] or "").strip() and parse_money(row[index]) is None),
            None,
        )
        if text is not None:
            shifted["description"] = text
    return shifted


def _extend(operation: RawOperation, tail: str) -> None:
    text = f"{operation.description} {tail}".strip()[:DESCRIPTION_LIMIT]
    operation.description = text
    operation.merchant = text[:200]


def _row_amount(row: list[Any], mapping: dict[str, int]) -> float | None:
    """Resolve a signed amount from either one column or a debit/credit pair."""
    amount = parse_money(_cell(row, mapping, "amount"))

    if amount is None:
        credit = parse_money(_cell(row, mapping, "credit")) or 0.0
        debit = parse_money(_cell(row, mapping, "debit")) or 0.0
        if not credit and not debit:
            return None
        return abs(credit) - abs(debit)

    # Some exports keep the amount unsigned and state the direction separately.
    direction = _text(row, mapping, "direction").lower().replace("ё", "е")
    if direction:
        if any(marker in direction for marker in DEBIT_MARKERS):
            return -abs(amount)
        if any(marker in direction for marker in CREDIT_MARKERS):
            return abs(amount)
    return amount
